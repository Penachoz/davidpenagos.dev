#!/usr/bin/env python3
"""Generate PDF and DOCX copies of the LinkedIn and cloud modernization guides."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate

from generate_dhl_application import markdown_flowables

ROOT = Path(__file__).resolve().parents[1]
ARTIFACTS = Path("/opt/cursor/artifacts")

DOCUMENTS = [
    (
        ROOT / "docs" / "linkedin_profile_refresh.md",
        "LinkedIn_David_Penagos_Optimizacion_2026",
        "Optimización de LinkedIn — David Penagos",
    ),
    (
        ROOT / "docs" / "api_cloud_modernization.md",
        "Plan_3_Proyectos_Cloud_y_APIs_David_Penagos",
        "Modernización de APIs y proyectos cloud — David Penagos",
    ),
]


def page_footer(canvas, doc) -> None:
    canvas.saveState()
    canvas.setStrokeColor(colors.HexColor("#D5DEE7"))
    canvas.setLineWidth(0.4)
    canvas.line(18 * mm, 10 * mm, A4[0] - 18 * mm, 10 * mm)
    canvas.setFont("DejaVu", 6.5)
    canvas.setFillColor(colors.HexColor("#667085"))
    canvas.drawString(18 * mm, 6.8 * mm, "DAVID PENAGOS · INFRASTRUCTURE / CLOUD / DEVOPS")
    canvas.drawRightString(A4[0] - 18 * mm, 6.8 * mm, str(canvas.getPageNumber()))
    canvas.restoreState()


def build_pdf(source: Path, destination: Path, title: str) -> None:
    doc = SimpleDocTemplate(
        str(destination),
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=15 * mm,
        title=title,
        author="David Penagos",
    )
    story = markdown_flowables(source.read_text(encoding="utf-8"))
    doc.build(story, onFirstPage=page_footer, onLaterPages=page_footer)


def add_inline(paragraph, text: str) -> None:
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
        else:
            paragraph.add_run(part)


def build_docx(source: Path, destination: Path, title: str) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    document.core_properties.title = title
    document.core_properties.author = "David Penagos"

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(4)

    lines = source.read_text(encoding="utf-8").splitlines()
    index = 0
    in_code = False
    while index < len(lines):
        stripped = lines[index].strip()
        if stripped.startswith("```"):
            in_code = not in_code
            index += 1
            continue
        if in_code:
            paragraph = document.add_paragraph()
            run = paragraph.add_run(lines[index])
            run.font.name = "Consolas"
            run.font.size = Pt(8)
            index += 1
            continue
        if not stripped or stripped == "---":
            index += 1
            continue
        if stripped.startswith("|"):
            table_lines: list[str] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            rows = [
                [cell.strip() for cell in row.strip("|").split("|")]
                for row in table_lines
                if not all(
                    re.fullmatch(r":?-{3,}:?", cell.strip())
                    for cell in row.strip("|").split("|")
                )
            ]
            if rows:
                table = document.add_table(rows=0, cols=len(rows[0]))
                table.style = "Table Grid"
                for row_index, cells in enumerate(rows):
                    row = table.add_row()
                    for cell_index, value in enumerate(cells):
                        add_inline(row.cells[cell_index].paragraphs[0], value)
                        if row_index == 0:
                            for run in row.cells[cell_index].paragraphs[0].runs:
                                run.bold = True
            continue
        if stripped.startswith("# "):
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(stripped[2:])
            run.bold = True
            run.font.name = "Arial"
            run.font.size = Pt(18)
            run.font.color.rgb = RGBColor(16, 29, 49)
        elif stripped.startswith("## "):
            paragraph = document.add_heading(level=1)
            add_inline(paragraph, stripped[3:])
        elif stripped.startswith("### "):
            paragraph = document.add_heading(level=2)
            add_inline(paragraph, stripped[4:])
        elif re.match(r"^-\s+", stripped):
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, stripped[2:])
        elif re.match(r"^\d+\.\s+", stripped):
            paragraph = document.add_paragraph(style="List Number")
            add_inline(paragraph, re.sub(r"^\d+\.\s+", "", stripped))
        elif stripped.startswith(">"):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.25)
            add_inline(paragraph, stripped.lstrip("> ").strip())
        else:
            paragraph = document.add_paragraph()
            add_inline(paragraph, stripped)
        index += 1

    document.save(destination)


def main() -> None:
    ARTIFACTS.mkdir(parents=True, exist_ok=True)
    for source, stem, title in DOCUMENTS:
        pdf = ARTIFACTS / f"{stem}.pdf"
        docx = ARTIFACTS / f"{stem}.docx"
        build_pdf(source, pdf, title)
        build_docx(source, docx, title)
        print(pdf)
        print(docx)


if __name__ == "__main__":
    main()

