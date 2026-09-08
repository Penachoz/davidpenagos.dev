#!/usr/bin/env python3
"""Generate an ATS resume (PDF + DOCX) and a DHL interview guide (PDF)."""

from __future__ import annotations

import html
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    HRFlowable,
    PageBreak,
    PageTemplate,
    Paragraph,
    Preformatted,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

ROOT = Path(__file__).resolve().parents[1]
RESUME_SOURCE = ROOT / "docs" / "dhl_cv_ats.md"
GUIDE_SOURCE = ROOT / "docs" / "dhl_guia_entrevista_coordinador_it.md"
ARTIFACTS = Path("/opt/cursor/artifacts")
RESUME_PDF = ARTIFACTS / "CV_David_Penagos_DHL_Coordinador_IT_ATS.pdf"
RESUME_DOCX = ARTIFACTS / "CV_David_Penagos_DHL_Coordinador_IT_ATS.docx"
GUIDE_PDF = ARTIFACTS / "Guia_entrevista_DHL_Coordinador_IT_David_Penagos.pdf"

NAVY = colors.HexColor("#101D31")
INK = colors.HexColor("#1B2638")
BLUE = colors.HexColor("#145A86")
YELLOW = colors.HexColor("#FFCC00")
PALE = colors.HexColor("#EEF5F8")
LIGHT = colors.HexColor("#F6F8FA")
MID = colors.HexColor("#667085")
LINE = colors.HexColor("#D5DEE7")
WHITE = colors.white

pdfmetrics.registerFont(
    TTFont("DejaVu", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
)
pdfmetrics.registerFont(
    TTFont("DejaVu-Bold", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf")
)
pdfmetrics.registerFont(
    TTFont("DejaVu-Mono", "/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf")
)


def inline_markup(text: str) -> str:
    tokens: list[str] = []

    def stash(value: str) -> str:
        tokens.append(value)
        return f"@@T{len(tokens) - 1}@@"

    text = re.sub(
        r"\[([^\]]+)\]\((https?://[^)]+)\)",
        lambda m: stash(
            f'<link href="{html.escape(m.group(2), quote=True)}" '
            f'color="#145A86"><u>{html.escape(m.group(1))}</u></link>'
        ),
        text,
    )
    text = re.sub(
        r"`([^`]+)`",
        lambda m: stash(
            f'<font name="DejaVu-Mono" color="#0B7182">{html.escape(m.group(1))}</font>'
        ),
        text,
    )
    text = html.escape(text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"<b>\1</b>", text)
    for idx, token in enumerate(tokens):
        text = text.replace(f"@@T{idx}@@", token)
    return text


base = getSampleStyleSheet()
BODY = ParagraphStyle(
    "BodyDHL",
    fontName="DejaVu",
    fontSize=9.2,
    leading=13.5,
    textColor=INK,
    spaceAfter=4,
)
BULLET = ParagraphStyle(
    "BulletDHL",
    parent=BODY,
    leftIndent=12,
    firstLineIndent=-7,
    bulletIndent=3,
    spaceAfter=2.5,
)
H1 = ParagraphStyle(
    "H1DHL",
    fontName="DejaVu-Bold",
    fontSize=17,
    leading=21,
    textColor=NAVY,
    spaceBefore=8,
    spaceAfter=7,
    keepWithNext=True,
)
H2 = ParagraphStyle(
    "H2DHL",
    fontName="DejaVu-Bold",
    fontSize=13,
    leading=17,
    textColor=BLUE,
    spaceBefore=7,
    spaceAfter=4,
    keepWithNext=True,
)
H3 = ParagraphStyle(
    "H3DHL",
    fontName="DejaVu-Bold",
    fontSize=10.5,
    leading=14,
    textColor=INK,
    spaceBefore=4,
    spaceAfter=3,
    keepWithNext=True,
)
CODE = ParagraphStyle(
    "CodeDHL",
    fontName="DejaVu-Mono",
    fontSize=7.6,
    leading=11,
    textColor=INK,
    leftIndent=7,
    rightIndent=7,
    borderColor=LINE,
    borderWidth=0.5,
    borderPadding=6,
    backColor=LIGHT,
    spaceBefore=3,
    spaceAfter=7,
)


def parse_table(lines: list[str]) -> Table:
    rows = []
    for line in lines:
        cells = [x.strip() for x in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            continue
        rows.append(
            [
                Paragraph(
                    inline_markup(cell),
                    ParagraphStyle(
                        f"cell-{len(rows)}",
                        parent=BODY,
                        fontSize=8,
                        leading=11,
                        spaceAfter=0,
                    ),
                )
                for cell in cells
            ]
        )
    table = Table(rows, repeatRows=1, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), NAVY),
                ("TEXTCOLOR", (0, 0), (-1, 0), WHITE),
                ("GRID", (0, 0), (-1, -1), 0.4, LINE),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [WHITE, LIGHT]),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    return table


def markdown_flowables(
    text: str,
    *,
    skip_until: str | None = None,
    resume: bool = False,
) -> list:
    lines = text.splitlines()
    idx = 0
    if skip_until:
        while idx < len(lines) and not lines[idx].startswith(skip_until):
            idx += 1
    story: list = []
    buffer: list[str] = []

    def flush():
        if buffer:
            story.append(Paragraph(inline_markup(" ".join(buffer)), BODY))
            buffer.clear()

    while idx < len(lines):
        line = lines[idx].rstrip()
        stripped = line.strip()
        if not stripped:
            flush()
            idx += 1
            continue
        if stripped == "---":
            flush()
            story.append(
                HRFlowable(
                    width="100%",
                    thickness=0.5,
                    color=LINE,
                    spaceBefore=4,
                    spaceAfter=6,
                )
            )
            idx += 1
            continue
        if stripped.startswith("```"):
            flush()
            idx += 1
            code_lines = []
            while idx < len(lines) and not lines[idx].strip().startswith("```"):
                code_lines.append(lines[idx])
                idx += 1
            idx += 1
            story.append(Preformatted("\n".join(code_lines), CODE))
            continue
        if stripped.startswith("|"):
            flush()
            table_lines = []
            while idx < len(lines) and lines[idx].strip().startswith("|"):
                table_lines.append(lines[idx])
                idx += 1
            story.extend([parse_table(table_lines), Spacer(1, 5)])
            continue
        if stripped.startswith("# "):
            flush()
            style = ParagraphStyle(
                "ResumeName" if resume else "DocumentTitle",
                fontName="DejaVu-Bold",
                fontSize=18 if resume else 17,
                leading=22,
                textColor=NAVY,
                alignment=TA_CENTER if resume else 0,
                spaceAfter=5,
            )
            story.append(Paragraph(inline_markup(stripped[2:]), style))
            idx += 1
            continue
        if stripped.startswith("## "):
            flush()
            if resume and len(story) < 2:
                style = ParagraphStyle(
                    "ResumeRole",
                    fontName="DejaVu-Bold",
                    fontSize=10.5,
                    leading=14,
                    textColor=BLUE,
                    alignment=TA_CENTER,
                    spaceAfter=5,
                )
            else:
                style = H2
            story.append(Paragraph(inline_markup(stripped[3:]), style))
            idx += 1
            continue
        if stripped.startswith("### "):
            flush()
            story.append(Paragraph(inline_markup(stripped[4:]), H3))
            idx += 1
            continue
        bullet = re.match(r"^-\s+(.+)$", stripped)
        numbered = re.match(r"^(\d+)\.\s+(.+)$", stripped)
        if bullet or numbered:
            flush()
            marker = "•" if bullet else f"{numbered.group(1)}."
            body = bullet.group(1) if bullet else numbered.group(2)
            story.append(
                Paragraph(f"<b>{marker}</b> {inline_markup(body)}", BULLET)
            )
            idx += 1
            continue
        if stripped.startswith(">"):
            flush()
            quote = stripped.lstrip("> ").strip()
            quote_lines = [quote]
            idx += 1
            while idx < len(lines) and lines[idx].strip().startswith(">"):
                quote_lines.append(lines[idx].strip().lstrip("> ").strip())
                idx += 1
            story.append(
                Paragraph(
                    inline_markup(" ".join(quote_lines)),
                    ParagraphStyle(
                        "QuoteDHL",
                        parent=BODY,
                        leftIndent=10,
                        rightIndent=7,
                        borderColor=BLUE,
                        borderWidth=1.2,
                        borderPadding=7,
                        backColor=PALE,
                        spaceBefore=3,
                        spaceAfter=7,
                    ),
                )
            )
            continue
        buffer.append(stripped)
        idx += 1
    flush()
    return story


def resume_page(canvas, doc):
    canvas.saveState()
    canvas.setStrokeColor(LINE)
    canvas.line(16 * mm, 11 * mm, A4[0] - 16 * mm, 11 * mm)
    canvas.setFont("DejaVu", 6.5)
    canvas.setFillColor(MID)
    canvas.drawString(16 * mm, 7.5 * mm, "David Alejandro Penagos Valencia · CV dirigido a DHL")
    canvas.drawRightString(A4[0] - 16 * mm, 7.5 * mm, str(canvas.getPageNumber()))
    canvas.restoreState()


class GuideDocument(BaseDocTemplate):
    def __init__(self, filename: str):
        super().__init__(
            filename,
            pagesize=A4,
            leftMargin=18 * mm,
            rightMargin=18 * mm,
            topMargin=17 * mm,
            bottomMargin=17 * mm,
            title="Guía de entrevista Coordinador IT ISM — DHL",
            author="David Alejandro Penagos Valencia",
        )
        self.addPageTemplates(
            [
                PageTemplate(
                    id="guide",
                    frames=[
                        Frame(
                            self.leftMargin,
                            self.bottomMargin,
                            self.width,
                            self.height,
                            id="guide-frame",
                        )
                    ],
                    onPage=self.page,
                )
            ]
        )

    @staticmethod
    def page(canvas, doc):
        page = canvas.getPageNumber()
        canvas.saveState()
        if page == 1:
            canvas.setFillColor(NAVY)
            canvas.rect(0, 0, A4[0], A4[1], fill=1, stroke=0)
            canvas.setFillColor(YELLOW)
            canvas.rect(0, A4[1] - 8 * mm, A4[0], 8 * mm, fill=1, stroke=0)
        else:
            canvas.setStrokeColor(LINE)
            canvas.line(
                18 * mm, A4[1] - 10 * mm, A4[0] - 18 * mm, A4[1] - 10 * mm
            )
            canvas.setFont("DejaVu", 6.5)
            canvas.setFillColor(MID)
            canvas.drawString(
                18 * mm,
                A4[1] - 7.5 * mm,
                "DHL · COORDINADOR IT–ISM · GUÍA DE PREPARACIÓN",
            )
            canvas.drawString(18 * mm, 8 * mm, "David Alejandro Penagos Valencia")
            canvas.drawRightString(
                A4[0] - 18 * mm, 8 * mm, str(canvas.getPageNumber())
            )
        canvas.restoreState()


def build_resume_pdf() -> None:
    doc = SimpleDocTemplate(
        str(RESUME_PDF),
        pagesize=A4,
        leftMargin=16 * mm,
        rightMargin=16 * mm,
        topMargin=12 * mm,
        bottomMargin=15 * mm,
        title="David Alejandro Penagos Valencia — CV DHL Coordinador IT",
        author="David Alejandro Penagos Valencia",
    )
    story = markdown_flowables(
        RESUME_SOURCE.read_text(encoding="utf-8"), resume=True
    )
    doc.build(story, onFirstPage=resume_page, onLaterPages=resume_page)


def add_docx_text(paragraph, text: str):
    """Add plain Markdown emphasis without compromising ATS extraction."""
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
        else:
            paragraph.add_run(part)


def build_resume_docx() -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    document.core_properties.title = (
        "David Alejandro Penagos Valencia — CV DHL Coordinador IT"
    )
    document.core_properties.author = "David Alejandro Penagos Valencia"

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.2)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.05

    lines = RESUME_SOURCE.read_text(encoding="utf-8").splitlines()
    for line in lines:
        stripped = line.strip()
        if not stripped or stripped == "---":
            continue
        if stripped.startswith("# "):
            p = document.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(stripped[2:])
            r.bold = True
            r.font.name = "Arial"
            r.font.size = Pt(17)
            r.font.color.rgb = RGBColor(16, 29, 49)
            p.paragraph_format.space_after = Pt(2)
            continue
        if stripped.startswith("## "):
            text = stripped[3:]
            if text.startswith("SOPORTE"):
                p = document.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(text)
                r.bold = True
                r.font.name = "Arial"
                r.font.size = Pt(10)
                r.font.color.rgb = RGBColor(20, 90, 134)
            else:
                p = document.add_paragraph()
                p.paragraph_format.space_before = Pt(5)
                p.paragraph_format.space_after = Pt(2)
                r = p.add_run(text)
                r.bold = True
                r.font.name = "Arial"
                r.font.size = Pt(10.5)
                r.font.color.rgb = RGBColor(20, 90, 134)
                bottom = OxmlElement("w:pBdr")
                border = OxmlElement("w:bottom")
                border.set(qn("w:val"), "single")
                border.set(qn("w:sz"), "4")
                border.set(qn("w:color"), "D5DEE7")
                bottom.append(border)
                p._p.get_or_add_pPr().append(bottom)
            continue
        if stripped.startswith("### "):
            p = document.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(stripped[4:])
            r.bold = True
            r.font.name = "Arial"
            r.font.size = Pt(9.5)
            continue
        if stripped.startswith("- "):
            p = document.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Inches(0.18)
            p.paragraph_format.first_line_indent = Inches(-0.12)
            p.paragraph_format.space_after = Pt(1.5)
            add_docx_text(p, stripped[2:])
            continue
        p = document.add_paragraph()
        if stripped.startswith(("Cali, Colombia", "LinkedIn:", "GitHub:")):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(1)
        add_docx_text(p, stripped)

    document.save(RESUME_DOCX)


def build_guide_pdf() -> None:
    doc = GuideDocument(str(GUIDE_PDF))
    cover = [
        Spacer(1, 45 * mm),
        Paragraph(
            "DHL · CALI · COORDINADOR IT–ISM",
            ParagraphStyle(
                "GuideEye",
                fontName="DejaVu-Bold",
                fontSize=10,
                leading=13,
                textColor=YELLOW,
                alignment=TA_CENTER,
            ),
        ),
        Spacer(1, 9),
        Paragraph(
            "Guía técnica y de entrevista",
            ParagraphStyle(
                "GuideCoverTitle",
                fontName="DejaVu-Bold",
                fontSize=27,
                leading=32,
                textColor=WHITE,
                alignment=TA_CENTER,
            ),
        ),
        Spacer(1, 10),
        Paragraph(
            "Redes · ITSM · Backups · Activos · KPI · SD-WAN",
            ParagraphStyle(
                "GuideCoverSub",
                fontName="DejaVu",
                fontSize=12,
                leading=17,
                textColor=colors.HexColor("#D9E3EC"),
                alignment=TA_CENTER,
            ),
        ),
        Spacer(1, 20 * mm),
        HRFlowable(
            width="45%",
            thickness=2,
            color=YELLOW,
            spaceBefore=4,
            spaceAfter=17,
            hAlign="CENTER",
        ),
        Paragraph(
            "David Alejandro Penagos Valencia",
            ParagraphStyle(
                "GuideName",
                fontName="DejaVu-Bold",
                fontSize=13,
                textColor=WHITE,
                alignment=TA_CENTER,
            ),
        ),
        Spacer(1, 7),
        Paragraph(
            "Postulación de alcance · preparación honesta · evidencia real",
            ParagraphStyle(
                "GuideTag",
                fontName="DejaVu",
                fontSize=9,
                textColor=colors.HexColor("#AAC0D1"),
                alignment=TA_CENTER,
            ),
        ),
        PageBreak(),
    ]
    body = markdown_flowables(
        GUIDE_SOURCE.read_text(encoding="utf-8"), skip_until="# 1."
    )
    doc.build(cover + body)


def main() -> None:
    ARTIFACTS.mkdir(parents=True, exist_ok=True)
    build_resume_pdf()
    build_resume_docx()
    build_guide_pdf()
    for path in [RESUME_PDF, RESUME_DOCX, GUIDE_PDF]:
        print(path)


if __name__ == "__main__":
    main()
